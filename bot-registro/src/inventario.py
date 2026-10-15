"""Fase 1: inventario de los ayuntamientos de la provincia de Valencia a partir del DIR3.

Pasos:
 1. Descargar el listado de unidades DIR3 de entidades locales y guardarlo sin
    tocar en datos/dir3_bruto/ con la fecha de descarga (o usar --archivo).
 2. Filtrar los ayuntamientos: código L0146xxxx.
 3. Excluir entidades locales menores, mancomunidades y organismos dependientes.
 4. Comprobar el recuento (266). Si no cuadra, se para y lo muestra.
 5. Guardar las unidades hijas de cada ayuntamiento.
 6. Opcional: cruzar con la población del INE.
 7. Crear o actualizar datos/ayuntamientos.xlsx sin borrar el estado de envíos.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook

from comun import (
    PATRON_AYTO_VALENCIA,
    RAIZ,
    elegir_unidad,
    municipio_desde_denominacion,
    normalizar,
)
from hoja import ENVIADO, PENDIENTE, HojaControl

CARPETA_BRUTO = RAIZ / "datos" / "dir3_bruto"

# Nombres de columna conocidos en las descargas del DIR3 (normalizados con _clave()).
ALIAS = {
    "codigo": [
        "c_id_ud_organica", "codigo_unidad", "codigo_de_la_unidad", "codigo_dir3",
        "cod_unidad", "codigo_unidad_organica", "id_unidad", "codigo", "cod_dir3", "dir3",
    ],
    "nombre": [
        "c_dnm_ud_organica", "denominacion", "denominacion_unidad", "nombre_unidad",
        "denominacion_de_la_unidad", "nombre", "unidad",
    ],
    "superior": [
        "c_id_ud_superior", "codigo_unidad_superior", "cod_unidad_superior",
        "codigo_de_la_unidad_superior", "unidad_superior", "codigo_superior",
    ],
    "raiz": [
        "c_id_ud_principal", "codigo_unidad_raiz", "cod_unidad_raiz", "codigo_de_la_unidad_raiz",
        "unidad_raiz", "codigo_raiz", "c_id_ud_raiz",
    ],
    "estado": ["c_id_estado", "estado", "estado_unidad", "cod_estado"],
    "web": ["c_url", "url", "web", "pagina_web", "url_web", "sitio_web", "direccion_web"],
}

ESTADOS_VIGENTES = {"v", "vigente", "activa", "activo"}

# Denominaciones que no son un ayuntamiento aunque el código lo parezca.
PATRON_NO_AYUNTAMIENTO = re.compile(
    r"mancomunitat|mancomunidad|entidad local menor|entitat local menor|consorci|"
    r"organismo autonomo|organisme autonom|fundacio|fundacion|empresa|sociedad|"
    r"patronato|patronat|instituto|institut",
    re.IGNORECASE,
)


class InventarioError(Exception):
    pass


def _clave(texto) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalizar(str(texto or ""))).strip("_")


# ---------------------------------------------------------------------------
# 1. Descarga
# ---------------------------------------------------------------------------

def descargar(url: str, log) -> Path:
    import requests

    CARPETA_BRUTO.mkdir(parents=True, exist_ok=True)
    log.info(f"Descargando DIR3 desde {url}")
    respuesta = requests.get(url, timeout=120, headers={"User-Agent": "Mozilla/5.0 bot-registro"})
    respuesta.raise_for_status()
    nombre = Path(unquote(urlparse(url).path)).name or "dir3"
    disposicion = respuesta.headers.get("content-disposition", "")
    m = re.search(r'filename="?([^";]+)"?', disposicion)
    if m:
        nombre = m.group(1)
    destino = CARPETA_BRUTO / f"{datetime.now():%Y-%m-%d}_{nombre}"
    destino.write_bytes(respuesta.content)
    log.info(f"Guardado sin modificar en {destino}")
    return destino


# ---------------------------------------------------------------------------
# Lectura de la tabla (CSV, XLSX o ZIP con uno de ellos)
# ---------------------------------------------------------------------------

def _filas_csv(contenido: bytes) -> list[list[str]]:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:5000]
    delimitador = max([";", ",", "\t", "|"], key=muestra.count)
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _filas_xlsx(contenido: bytes) -> list[list[str]]:
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    mejor: list[list[str]] = []
    for hoja in libro.worksheets:
        filas = [["" if v is None else str(v).strip() for v in fila] for fila in hoja.iter_rows(values_only=True)]
        if len(filas) > len(mejor):
            mejor = filas
    return mejor


def leer_filas_brutas(ruta: Path) -> list[list[str]]:
    sufijo = ruta.suffix.lower()
    contenido = ruta.read_bytes()
    if sufijo == ".zip":
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            candidatos = [n for n in z.namelist() if n.lower().endswith((".csv", ".xlsx", ".txt"))]
            if not candidatos:
                raise InventarioError(f"El ZIP {ruta.name} no contiene CSV ni XLSX")
            # el mayor suele ser el listado de unidades
            nombre = max(candidatos, key=lambda n: z.getinfo(n).file_size)
            datos = z.read(nombre)
            return _filas_xlsx(datos) if nombre.lower().endswith(".xlsx") else _filas_csv(datos)
    if sufijo in (".xlsx", ".xlsm"):
        return _filas_xlsx(contenido)
    if sufijo in (".csv", ".txt"):
        return _filas_csv(contenido)
    raise InventarioError(
        f"Formato {sufijo} no soportado. Abre el archivo en Excel y guárdalo como .xlsx o .csv"
    )


def detectar_columnas(cabecera: list[str], forzadas: dict | None = None) -> dict[str, int]:
    claves = [_clave(c) for c in cabecera]
    columnas: dict[str, int] = {}
    for campo, nombre in (forzadas or {}).items():
        if _clave(nombre) in claves:
            columnas[campo] = claves.index(_clave(nombre))
    for campo, alias in ALIAS.items():
        if campo in columnas:
            continue
        for a in alias:
            if a in claves:
                columnas[campo] = claves.index(a)
                break
    return columnas


def leer_unidades(ruta: Path, forzadas: dict | None = None) -> list[dict]:
    """Devuelve una lista de unidades {codigo, nombre, superior, raiz, estado, web}."""
    filas = leer_filas_brutas(ruta)
    # La cabecera es la primera fila (de las 20 primeras) en la que se reconoce código y nombre.
    for i, fila in enumerate(filas[:20]):
        columnas = detectar_columnas(fila, forzadas)
        if "codigo" in columnas and "nombre" in columnas:
            break
    else:
        cabeceras = filas[0] if filas else []
        raise InventarioError(
            "No se reconocen las columnas de código y denominación en el DIR3. "
            f"Cabecera encontrada: {cabeceras}. Indica los nombres en config.yaml, "
            "inventario.columnas (codigo, nombre, superior, raiz, estado, web)."
        )
    unidades = []
    for fila in filas[i + 1:]:
        def valor(campo):
            j = columnas.get(campo)
            return fila[j].strip() if j is not None and j < len(fila) and fila[j] else ""
        codigo = valor("codigo").upper()
        if not codigo:
            continue
        unidades.append({
            "codigo": codigo,
            "nombre": valor("nombre"),
            "superior": valor("superior").upper(),
            "raiz": valor("raiz").upper(),
            "estado": valor("estado"),
            "web": valor("web"),
        })
    return unidades


# ---------------------------------------------------------------------------
# 2-5. Filtrado y unidades hijas
# ---------------------------------------------------------------------------

def solo_vigentes(unidades: list[dict], log) -> list[dict]:
    con_estado = [u for u in unidades if u["estado"]]
    if not con_estado:
        return unidades
    vigentes = [u for u in unidades if normalizar(u["estado"]) in ESTADOS_VIGENTES]
    if not vigentes:
        valores = sorted({u["estado"] for u in con_estado})[:10]
        log.warning(f"No se reconoce ningún estado vigente (valores: {valores}); no se filtra por estado")
        return unidades
    return vigentes


def filtrar_ayuntamientos(unidades: list[dict]) -> tuple[list[dict], list[dict]]:
    """Devuelve (ayuntamientos, descartados_con_codigo_L0146)."""
    ayuntamientos, descartados = [], []
    for u in unidades:
        if not PATRON_AYTO_VALENCIA.match(u["codigo"]):
            continue
        if PATRON_NO_AYUNTAMIENTO.search(normalizar(u["nombre"])):
            descartados.append(u)
        else:
            ayuntamientos.append(u)
    return ayuntamientos, descartados


def unidades_hijas(ayuntamientos: list[dict], unidades: list[dict]) -> dict[str, list[dict]]:
    """Unidades que cuelgan (directa o indirectamente) de cada ayuntamiento."""
    codigos_ayto = {a["codigo"] for a in ayuntamientos}
    superior = {u["codigo"]: u["superior"] for u in unidades}
    hijas: dict[str, list[dict]] = {c: [] for c in codigos_ayto}
    for u in unidades:
        if u["codigo"] in codigos_ayto:
            continue
        ayto = u["raiz"] if u["raiz"] in codigos_ayto else ""
        if not ayto:
            actual, pasos = u["superior"], 0
            while actual and pasos < 15:
                if actual in codigos_ayto:
                    ayto = actual
                    break
                actual, pasos = superior.get(actual, ""), pasos + 1
        if ayto:
            hijas[ayto].append(u)
    return hijas


# ---------------------------------------------------------------------------
# 6. Población del INE (opcional)
# ---------------------------------------------------------------------------

def leer_poblacion(ruta: Path) -> dict[str, int]:
    """Lee el fichero de población del INE (pobmunXX.xlsx) o un CSV codigo_ine;poblacion.

    Devuelve {codigo_ine_5_digitos: habitantes}.
    """
    filas = leer_filas_brutas(ruta)
    for i, fila in enumerate(filas[:20]):
        claves = [_clave(c) for c in fila]
        if ("cpro" in claves and "cmun" in claves) or "codigo_ine" in claves:
            break
    else:
        raise InventarioError("No se reconocen las columnas del fichero de población (CPRO/CMUN o codigo_ine)")
    col_pob = next((j for j, c in enumerate(claves) if c.startswith("pob")), None)
    if col_pob is None:
        raise InventarioError("No se encuentra la columna de población (POBxx o poblacion)")
    resultado = {}
    for fila in filas[i + 1:]:
        try:
            if "codigo_ine" in claves:
                codigo = re.sub(r"\D", "", fila[claves.index("codigo_ine")])[:5].zfill(5)
            else:
                codigo = fila[claves.index("cpro")].zfill(2) + fila[claves.index("cmun")].zfill(3)
            resultado[codigo] = _entero(fila[col_pob])
        except (ValueError, IndexError):
            continue
    return resultado


def _entero(valor) -> int:
    texto = str(valor).strip()
    if re.fullmatch(r"\d+(\.0+)?", texto):
        return int(float(texto))
    # formato con separador de miles: 12.345 o 12,345
    return int(re.sub(r"[.,\s]", "", texto))


def codigo_ine(dir3_ayto: str) -> str:
    """L01462508 -> 46250."""
    return dir3_ayto[3:8]


def calcular_prioridad(poblacion, umbrales: dict) -> int:
    try:
        habitantes = int(poblacion)
    except (TypeError, ValueError):
        return 2
    if habitantes >= int(umbrales.get("umbral_1", 10000)):
        return 1
    if habitantes >= int(umbrales.get("umbral_2", 3000)):
        return 2
    return 3


# ---------------------------------------------------------------------------
# 7. Hoja de control
# ---------------------------------------------------------------------------

def actualizar_hoja(hoja: HojaControl, ayuntamientos, hijas, poblacion, umbrales, log) -> None:
    hoy = datetime.now().strftime("%Y-%m-%d")
    vistos = set()
    for a in sorted(ayuntamientos, key=lambda x: x["codigo"]):
        codigo = a["codigo"]
        vistos.add(codigo)
        fila = hoja.get(codigo) or hoja.nueva_fila(codigo)
        fila["municipio"] = municipio_desde_denominacion(a["nombre"])
        if a.get("web"):
            fila["web_dir3"] = a["web"]
        pob = poblacion.get(codigo_ine(codigo)) if poblacion else None
        if pob is not None:
            fila["poblacion"] = pob
        if not fila.get("prioridad"):
            fila["prioridad"] = calcular_prioridad(fila.get("poblacion"), umbrales)
        lista = [{"dir3_ayto": codigo, "codigo": h["codigo"], "nombre": h["nombre"],
                  "codigo_superior": h["superior"]} for h in hijas.get(codigo, [])]
        hoja.unidades[codigo] = lista
        # Sugerencia de destinatario: solo si todavía no se ha enviado ni fijado.
        if fila.get("estado") == PENDIENTE and not fila.get("dir3_destino"):
            destino, nombre = elegir_unidad(codigo, a["nombre"], [(u["codigo"], u["nombre"]) for u in lista])
            fila["dir3_destino"], fila["nombre_destino"] = destino, nombre
    for codigo, fila in hoja.filas.items():
        if codigo not in vistos:
            aviso = f"No figura en el DIR3 de {hoy}"
            if aviso not in str(fila.get("notas", "")):
                fila["notas"] = f"{fila.get('notas', '')} | {aviso}".strip(" |")
            log.warning(f"{codigo} {fila.get('municipio')}: {aviso} (se conserva la fila)")


def ejecutar(config: dict, log, archivo: str | None = None, poblacion: str | None = None,
             forzar: bool = False) -> HojaControl:
    cfg = config["inventario"]
    if archivo:
        ruta = Path(archivo)
        if not ruta.exists():
            raise InventarioError(f"No existe {ruta}")
        CARPETA_BRUTO.mkdir(parents=True, exist_ok=True)
        if CARPETA_BRUTO.resolve() not in ruta.resolve().parents:
            copia = CARPETA_BRUTO / f"{datetime.now():%Y-%m-%d}_{ruta.name}"
            copia.write_bytes(ruta.read_bytes())
            log.info(f"Copia sin modificar en {copia}")
    elif cfg.get("url_dir3"):
        ruta = descargar(cfg["url_dir3"], log)
    else:
        raise InventarioError(
            "Falta el listado DIR3. Pon la URL de descarga en config.yaml (inventario.url_dir3) "
            "o descárgalo a mano y ejecuta: python src/main.py inventario --archivo <ruta>"
        )

    unidades = solo_vigentes(leer_unidades(ruta, cfg.get("columnas")), log)
    log.info(f"Unidades vigentes leídas: {len(unidades)}")
    ayuntamientos, descartados = filtrar_ayuntamientos(unidades)
    for d in descartados:
        log.info(f"Excluido (no es ayuntamiento): {d['codigo']} {d['nombre']}")

    # códigos duplicados (no debería ocurrir)
    por_codigo: dict[str, list] = {}
    for a in ayuntamientos:
        por_codigo.setdefault(a["codigo"], []).append(a)
    duplicados = {c: l for c, l in por_codigo.items() if len(l) > 1}
    ayuntamientos = [l[0] for l in por_codigo.values()]

    datos_poblacion = {}
    ruta_pob = poblacion
    if not ruta_pob and cfg.get("url_poblacion"):
        ruta_pob = str(descargar(cfg["url_poblacion"], log))
    if ruta_pob:
        datos_poblacion = leer_poblacion(Path(ruta_pob))
        log.info(f"Población INE leída para {len(datos_poblacion)} municipios")

    esperado = int(cfg.get("total_esperado", 266))
    if len(ayuntamientos) != esperado or duplicados:
        log.error(f"Recuento: {len(ayuntamientos)} ayuntamientos; se esperaban {esperado}.")
        for c, l in duplicados.items():
            log.error(f"  Duplicado: {c} -> {[x['nombre'] for x in l]}")
        if datos_poblacion:
            ine_valencia = {c for c in datos_poblacion if c.startswith("46")}
            ine_dir3 = {codigo_ine(a["codigo"]) for a in ayuntamientos}
            for c in sorted(ine_valencia - ine_dir3):
                log.error(f"  Falta en DIR3 el municipio INE {c}")
            for c in sorted(ine_dir3 - ine_valencia):
                nombres = [a["nombre"] for a in ayuntamientos if codigo_ine(a["codigo"]) == c]
                log.error(f"  Sobra (no está en el INE) {c}: {nombres}")
        else:
            log.error("  Para ver cuáles faltan o sobran, añade --poblacion <pobmun.xlsx del INE>.")
            for a in sorted(ayuntamientos, key=lambda x: x["codigo"]):
                log.info(f"    {a['codigo']}  {a['nombre']}")
        if not forzar:
            raise InventarioError(
                "El recuento no cuadra; no se ha modificado la hoja. Revisa la lista y, "
                "si es correcta, repite con --forzar."
            )
        log.warning("Se continúa por --forzar")

    hijas = unidades_hijas(ayuntamientos, unidades)
    hoja = HojaControl()
    copia = hoja.copia_seguridad()
    if copia:
        log.info(f"Copia de la hoja en {copia}")
    actualizar_hoja(hoja, ayuntamientos, hijas, datos_poblacion, cfg.get("prioridad", {}), log)
    hoja.guardar()

    con_fiestas = sum(1 for f in hoja if f.get("dir3_destino") and f["dir3_destino"] != f["dir3_ayto"])
    enviados = sum(1 for f in hoja if f.get("estado") == ENVIADO)
    log.info(f"Hoja {hoja.ruta} actualizada: {len(hoja)} ayuntamientos, "
             f"{sum(len(v) for v in hijas.values())} unidades hijas, "
             f"{con_fiestas} con unidad de Fiestas/Cultura en DIR3, {enviados} ya enviados.")
    return hoja
