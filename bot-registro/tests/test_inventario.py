"""Lectura del DIR3 con el formato real del CTT ("Listado Unidades EELL.xlsx", octubre 2026)."""

from openpyxl import Workbook

from inventario import buscar_homonimos, filtrar_ayuntamientos, leer_unidades, solo_vigentes, unidades_hijas

CABECERA = [None, "C_ID_UD_ORGANICA", "C_DNM_UD_ORGANICA", "C_ID_NIVEL_ADMON", "C_ID_TIPO_ENT_PUBLICA",
            "N_NIVEL_JERARQUICO", "C_ID_AMB_PROVINCIA", "C_DESC_PROV", "C_ID_DEP_UD_SUPERIOR",
            "C_DNM_UD_ORGANICA", "C_ID_DEP_UD_PRINCIPAL", "C_DNM_UD_ORGANICA", "B_SW_DEP_EDP_PRINCIPAL",
            "C_ID_DEP_EDP_PRINCIPAL", "C_DNM_UD_ORGANICA", "C_ID_ESTADO", "D_VIG_ALTA_OFICIAL", "NIF_CIF",
            "CONTACTOS"]


def fila(codigo, nombre, tipo, superior, nombre_sup, raiz, estado="V", email=None):
    return [None, codigo, nombre, "3", tipo, "1", "46", "València/Valencia", superior, nombre_sup, raiz,
            nombre_sup, "N", "", "", estado, "", "P4600000A", email]


def test_formato_real_ctt(tmp_path):
    libro = Workbook()
    leyenda = libro.active
    leyenda.title = "Leyenda de campos"  # hoja de leyenda antes que la de datos
    leyenda.append([None, "C_ID_UD_ORGANICA", "Código de la UNIDAD orgánica", "Código que permite...", None])
    leyenda.append([None, "C_DNM_UD_ORGANICA", "Nombre de la UNIDAD orgánica", "Denominación", None])
    datos = libro.create_sheet("Unidades EELL")
    datos.append(CABECERA)
    datos.append(fila("LA9999999", "Entidades Locales", "", "LA9999999", "Entidades Locales", "LA9999999"))
    datos.append(fila("LA0008665", "Consorcio para la Gestión de Residuos Zona V5", "CS", "LA9999999",
                      "Entidades Locales", "LA0008665"))
    datos.append(fila("L01462508", "Ayuntamiento de València", "AY", "LA9999999", "Entidades Locales",
                      "L01462508", email="info@valencia.es"))
    datos.append(fila("LA0001111", "Concejalía de Fiestas", "AY", "L01462508", "Ayuntamiento de València",
                      "L01462508"))
    datos.append(fila("LA0001112", "Fundación Deportiva Municipal", "OA", "L01462508",
                      "Ayuntamiento de València", "L01462508"))
    datos.append(fila("L01460011", "Ayuntamiento de Albal", "AY", "LA9999999", "Entidades Locales", "L01460011"))
    datos.append(fila("L01460099", "Ayuntamiento de Extinto", "AY", "LA9999999", "Entidades Locales",
                      "L01460099", estado="E"))
    datos.append(fila("L01469998", "Algo con código de ayuntamiento", "MA", "LA9999999", "Entidades Locales",
                      "L01469998"))
    libro.create_sheet("Catálogos de clasificación").append([None, "Cód.", "Estado", None, None])
    ruta = tmp_path / "Listado Unidades EELL.xlsx"
    libro.save(ruta)

    unidades = solo_vigentes(leer_unidades(ruta), log=None)
    valencia = next(u for u in unidades if u["codigo"] == "L01462508")
    assert valencia["nombre"] == "Ayuntamiento de València"   # la primera C_DNM_UD_ORGANICA
    assert valencia["superior"] == "LA9999999"
    assert valencia["tipo"] == "AY" and valencia["email"] == "info@valencia.es"
    assert all(u["codigo"] != "L01460099" for u in unidades)    # extinguido

    ayuntamientos, descartados = filtrar_ayuntamientos(unidades)
    assert sorted(a["codigo"] for a in ayuntamientos) == ["L01460011", "L01462508"]
    assert [d["codigo"] for d in descartados] == ["L01469998"]  # tipo MA

    hijas = unidades_hijas(ayuntamientos, unidades)
    assert sorted(h["codigo"] for h in hijas["L01462508"]) == ["LA0001111", "LA0001112"]


def test_homonimos():
    unidades = [
        {"codigo": "L01460193", "nombre": "Ayuntamiento de Alcúdia, l'", "tipo": "AY"},
        {"codigo": "L01070034", "nombre": "Ajuntament d'Alcúdia", "tipo": "AY"},
        {"codigo": "L01460076", "nombre": "Ayuntamiento de Albal", "tipo": "AY"},
        {"codigo": "LA0000001", "nombre": "Ayuntamiento de Albal", "tipo": "AY"},  # no es un L01
    ]
    homonimos = buscar_homonimos(unidades[:1] + unidades[2:3], unidades)
    assert homonimos == {"L01460193": ["Alcúdia (L01070034)"]}
