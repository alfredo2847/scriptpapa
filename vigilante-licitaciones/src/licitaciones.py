"""Licitaciones abiertas: datos/licitaciones.xlsx.

La hoja se actualiza en cada búsqueda sin perder lo que escribes a mano en las
columnas `seguimiento` y `notas`.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from comun import DATOS, guardar_libro
from filtros import DUDOSO, ORDEN_ENCAJE, en_zona, municipio, puntuar
from placsp import ESTADOS, PROCEDIMIENTOS, TIPOS, Contrato

RUTA = DATOS / "licitaciones.xlsx"

NUEVA = "NUEVA"
ME_INTERESA = "ME INTERESA"
PRESENTADA = "PRESENTADA"
DESCARTADA = "DESCARTADA"
CERRADAS = (PRESENTADA, DESCARTADA)

COLUMNAS = [
    "seguimiento", "encaje", "dias_restantes", "plazo", "municipio", "organo", "objeto",
    "importe", "estado", "procedimiento", "tipo", "motivo", "enlace", "fuente",
    "expediente", "detectada", "actualizada", "notas", "id",
]
COLUMNAS_MANUALES = ("seguimiento", "notas")
ANCHOS = {"objeto": 60, "organo": 40, "motivo": 40, "enlace": 30, "notas": 30, "id": 20}


def abiertas(contratos: list[Contrato], config: dict, ahora: datetime) -> dict[str, dict]:
    """Licitaciones de la zona, con plazo vigente, que encajan."""
    resultado = {}
    for c in contratos:
        if c.estado not in ("PRE", "PUB"):
            continue
        if c.plazo and c.plazo < ahora:
            continue
        if not en_zona(c, config["provincias"]):
            continue
        puntuacion = puntuar(c, config)
        if not puntuacion:
            continue
        if puntuacion[0] == DUDOSO and not config["buscar"]["incluir_dudosos"]:
            continue
        resultado[c.id] = fila(c, puntuacion[0], puntuacion[1])
    return resultado


def fila(c: Contrato, encaje: str, motivo: str) -> dict:
    return {
        "encaje": encaje,
        "plazo": c.plazo,
        "municipio": municipio(c),
        "organo": c.organo,
        "objeto": c.objeto or c.titulo,
        "importe": c.importe,
        "estado": ESTADOS.get(c.estado, c.estado),
        "procedimiento": PROCEDIMIENTOS.get(c.procedimiento, c.procedimiento and f"código {c.procedimiento}"),
        "tipo": TIPOS.get(c.tipo, c.tipo),
        "motivo": motivo,
        "enlace": c.enlace,
        "fuente": c.fuente,
        "expediente": c.expediente,
        "actualizada": c.actualizado,
        "id": c.id,
    }


def cargar(ruta: Path | None = None) -> dict[str, dict]:
    ruta = ruta or RUTA
    if not ruta.exists():
        return {}
    hoja = load_workbook(ruta).active
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c or "") for c in next(filas, [])]
    existentes = {}
    for valores in filas:
        datos = dict(zip(cabecera, valores))
        if datos.get("id"):
            existentes[str(datos["id"])] = datos
    return existentes


def actualizar(existentes: dict[str, dict], encontradas: dict[str, dict], hoy: datetime) -> list[dict]:
    """Mezcla lo encontrado con la hoja. Devuelve las filas nuevas."""
    nuevas = []
    for ident, datos in encontradas.items():
        if ident in existentes:
            previa = existentes[ident]
            previa.update({k: v for k, v in datos.items() if k not in COLUMNAS_MANUALES})
        else:
            datos = {**datos, "seguimiento": NUEVA, "notas": "", "detectada": hoy.date()}
            existentes[ident] = datos
            nuevas.append(datos)
    return nuevas


def _clave_orden(f: dict, ahora: datetime):
    plazo = f.get("plazo") if isinstance(f.get("plazo"), datetime) else None
    vencida = bool(plazo and plazo < ahora)
    cerrada = f.get("seguimiento") in CERRADAS
    return (cerrada or vencida, ORDEN_ENCAJE.get(f.get("encaje"), 9), plazo or datetime.max)


def guardar(existentes: dict[str, dict], ahora: datetime, ruta: Path | None = None) -> None:
    ruta = ruta or RUTA
    libro = Workbook()
    hoja = libro.active
    hoja.title = "licitaciones"
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="DDEBF7")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    gris = Font(color="808080")
    for f in sorted(existentes.values(), key=lambda f: _clave_orden(f, ahora)):
        plazo = f.get("plazo") if isinstance(f.get("plazo"), datetime) else None
        f["dias_restantes"] = (plazo.date() - ahora.date()).days if plazo else None
        hoja.append([f.get(c) for c in COLUMNAS])
        n = hoja.max_row
        if f.get("seguimiento") == NUEVA:
            for celda in hoja[n]:
                celda.fill = amarillo
        elif f.get("seguimiento") in CERRADAS or (plazo and plazo < ahora):
            for celda in hoja[n]:
                celda.font = gris
    for i, columna in enumerate(COLUMNAS, start=1):
        letra = get_column_letter(i)
        hoja.column_dimensions[letra].width = ANCHOS.get(columna, max(len(columna) + 2, 12))
        formato = {"importe": '#,##0.00 "€"', "plazo": "dd/mm/yyyy hh:mm",
                   "actualizada": "dd/mm/yyyy hh:mm", "detectada": "dd/mm/yyyy"}.get(columna)
        if formato:
            for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = hoja.dimensions
    guardar_libro(libro, ruta)


def recordatorios(existentes: dict[str, dict], ahora: datetime, dias: int) -> list[dict]:
    """Las marcadas ME INTERESA cuyo plazo vence en `dias` días o menos."""
    salida = []
    for f in existentes.values():
        plazo = f.get("plazo") if isinstance(f.get("plazo"), datetime) else None
        if f.get("seguimiento") == ME_INTERESA and plazo and ahora <= plazo and (plazo.date() - ahora.date()).days <= dias:
            salida.append(f)
    return sorted(salida, key=lambda f: f["plazo"])
