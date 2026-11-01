"""Histórico: qué ayuntamientos contratan paellas, cuándo, por cuánto y a quién.

Genera datos/compradores.xlsx con tres pestañas:
- municipios: uno por fila, con prioridad sugerida y mes para enviar la propuesta;
- contratos: cada contrato encontrado;
- competencia: empresas adjudicatarias.

La hoja se rehace entera en cada ejecución (no tiene columnas para rellenar a mano).
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from comun import DATOS, MESES, guardar_libro
from filtros import ALTO, MEDIO, ORDEN_ENCAJE, dir3_ayuntamiento, en_zona, municipio, puntuar
from placsp import ESTADOS, Contrato

RUTA = DATOS / "compradores.xlsx"


@dataclass
class Hallazgo:
    contrato: Contrato
    encaje: str
    motivo: str
    municipio: str
    dir3: str

    @property
    def fecha(self) -> date | None:
        c = self.contrato
        if c.fecha_adjudicacion:
            return c.fecha_adjudicacion
        if c.plazo:
            return c.plazo.date()
        return c.actualizado.date() if c.actualizado else None

    @property
    def importe(self) -> float | None:
        c = self.contrato
        return c.importe_adjudicacion if c.importe_adjudicacion is not None else c.importe


def filtrar(contratos: Iterable[Contrato], config: dict) -> list[Hallazgo]:
    """Se queda con la última versión de cada contrato de la zona que encaja."""
    hallazgos: dict[str, Hallazgo] = {}
    for c in contratos:
        if c.estado == "ANUL" or not en_zona(c, config["provincias"]):
            continue
        resultado = puntuar(c, config)
        if not resultado:
            continue
        previo = hallazgos.get(c.id)
        if previo and (c.actualizado or datetime.min) <= (previo.contrato.actualizado or datetime.min):
            continue
        hallazgos[c.id] = Hallazgo(c, resultado[0], resultado[1], municipio(c),
                                   dir3_ayuntamiento(c, config["provincias"]))
    return list(hallazgos.values())


def _mes_envio(mes: int, antelacion: int) -> str:
    return MESES[(mes - 1 - antelacion) % 12]


def resumen_municipios(hallazgos: list[Hallazgo], antelacion: int = 2) -> list[dict]:
    grupos: dict[str, list[Hallazgo]] = defaultdict(list)
    for h in hallazgos:
        clave = h.dir3 or h.municipio.lower() or h.contrato.organo.lower()
        grupos[clave].append(h)

    filas = []
    for grupo in grupos.values():
        grupo.sort(key=lambda h: h.fecha or date.min, reverse=True)
        meses = Counter(h.fecha.month for h in grupo if h.fecha)
        mes_principal = meses.most_common(1)[0][0] if meses else None
        encajes = {h.encaje for h in grupo}
        prioridad = 1 if ALTO in encajes else 2 if MEDIO in encajes else 3
        ultimo = grupo[0]
        filas.append({
            "prioridad_sugerida": prioridad,
            "municipio": next((h.municipio for h in grupo if h.municipio), ""),
            "dir3_ayto": next((h.dir3 for h in grupo if h.dir3), ""),
            "contratos": len(grupo),
            "de_encaje_alto": sum(h.encaje == ALTO for h in grupo),
            "meses_en_que_contrata": ", ".join(f"{MESES[m - 1]} ({n})" for m, n in sorted(meses.items())),
            "enviar_propuesta_en": _mes_envio(mes_principal, antelacion) if mes_principal else "",
            "ultimo_contrato": ultimo.fecha,
            "ultimo_objeto": ultimo.contrato.objeto or ultimo.contrato.titulo,
            "ultimo_importe": ultimo.importe,
            "adjudicatarios": "; ".join(sorted({a for h in grupo for a in h.contrato.adjudicatarios})),
            "organos": "; ".join(sorted({h.contrato.organo for h in grupo if h.contrato.organo})),
        })
    filas.sort(key=lambda f: (f["prioridad_sugerida"], -f["contratos"], f["municipio"]))
    return filas


def resumen_competencia(hallazgos: list[Hallazgo]) -> list[dict]:
    grupos: dict[str, list[Hallazgo]] = defaultdict(list)
    for h in hallazgos:
        for a in h.contrato.adjudicatarios:
            grupos[a].append(h)
    filas = [{
        "empresa": empresa,
        "contratos": len(g),
        "importe_total": round(sum(h.importe or 0 for h in g), 2),
        "municipios": "; ".join(sorted({h.municipio for h in g if h.municipio})),
        "ultimo": max((h.fecha for h in g if h.fecha), default=None),
    } for empresa, g in grupos.items()]
    filas.sort(key=lambda f: (-f["contratos"], -f["importe_total"]))
    return filas


def filas_contratos(hallazgos: list[Hallazgo]) -> list[dict]:
    orden = sorted(hallazgos, key=lambda h: (ORDEN_ENCAJE[h.encaje], -(h.fecha or date.min).toordinal()))
    return [{
        "encaje": h.encaje,
        "fecha": h.fecha,
        "municipio": h.municipio,
        "organo": h.contrato.organo,
        "objeto": h.contrato.objeto or h.contrato.titulo,
        "importe": h.importe,
        "adjudicatario": "; ".join(h.contrato.adjudicatarios),
        "estado": ESTADOS.get(h.contrato.estado, h.contrato.estado),
        "motivo": h.motivo,
        "fuente": h.contrato.fuente,
        "dir3_ayto": h.dir3,
        "enlace": h.contrato.enlace,
        "id": h.contrato.id,
    } for h in orden]


def _pestana(libro: Workbook, titulo: str, filas: list[dict], primera: bool = False) -> None:
    hoja = libro.active if primera else libro.create_sheet()
    hoja.title = titulo
    if not filas:
        hoja.append(["(sin resultados)"])
        return
    columnas = list(filas[0])
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="DDEBF7")
    for fila in filas:
        hoja.append([fila[c] for c in columnas])
    for i, columna in enumerate(columnas, start=1):
        largo = max(len(str(columna)), *(min(len(str(f[columna] or "")), 60) for f in filas))
        hoja.column_dimensions[get_column_letter(i)].width = largo + 2
        if columna in ("importe", "ultimo_importe", "importe_total"):
            for celda in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                celda[0].number_format = '#,##0.00 "€"'
        if columna in ("fecha", "ultimo_contrato", "ultimo"):
            for celda in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                celda[0].number_format = "dd/mm/yyyy"
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions


def guardar(hallazgos: list[Hallazgo], antelacion: int, ruta: Path | None = None) -> dict:
    ruta = ruta or RUTA
    municipios = resumen_municipios(hallazgos, antelacion)
    libro = Workbook()
    _pestana(libro, "municipios", municipios, primera=True)
    _pestana(libro, "contratos", filas_contratos(hallazgos))
    _pestana(libro, "competencia", resumen_competencia(hallazgos))
    guardar_libro(libro, ruta)
    return {
        "contratos": len(hallazgos),
        "municipios": len(municipios),
        "prioridad_1": sum(f["prioridad_sugerida"] == 1 for f in municipios),
    }
