import zipfile
from datetime import date, datetime

from conftest import entrada, feed

import compradores
import licitaciones
from filtros import ALTO, DUDOSO, MEDIO, en_zona, municipio, puntuar
from placsp import leer_archivo, leer_atom, leer_feed_reciente, urls_historico

AHORA = datetime(2026, 11, 1, 9, 0)


def uno(**kw):
    kw.setdefault("objeto", "Servicio de paellas gigantes para las Fiestas Mayores")
    return next(leer_atom(feed(entrada("1", **kw))))


def test_parsea_campos():
    c = uno(lote="Lote 2: comida popular")
    assert c.id.endswith("/1")
    assert c.enlace.endswith("idEvl=1")
    assert c.estado == "PUB"
    assert c.organo == "Alcaldía del Ayuntamiento de Sueca"
    assert c.dir3 == ["LA0001234", "L01462357"]
    assert c.nif == "P4600000A"
    assert c.cp == "46410"
    assert c.importe == 12000
    assert c.cpvs == ["55520000"]
    assert c.nuts == ["ES523"]
    assert c.plazo == datetime(2026, 11, 20, 14, 0)
    assert c.procedimiento == "9"
    assert c.actualizado == datetime(2026, 10, 30, 9, 0)
    assert "comida popular" in c.texto
    assert municipio(c) == "Sueca"


def test_adjudicacion():
    c = uno(estado="RES", adjudicatario="Paellas Pepe SL", adjudicado="9500.5", fecha_adj="2025-07-02")
    assert c.adjudicatarios == ["Paellas Pepe SL"]
    assert c.importe_adjudicacion == 9500.5
    assert c.fecha_adjudicacion == date(2025, 7, 2)


def test_encaje(config):
    assert puntuar(uno(objeto="Servicio de paellas gigantes"), config)[0] == ALTO
    assert puntuar(uno(objeto="Dinar de germanor festes patronals"), config)[0] == ALTO
    assert puntuar(uno(objeto="Catering para las fiestas patronales", cpv="79952000"), config)[0] == MEDIO
    assert puntuar(uno(objeto="Servicio de comidas Fira de Tots Sants", cpv="90000000"), config)[0] == MEDIO
    assert puntuar(uno(objeto="Servicio de catering", cpv="55520000"), config)[0] == DUDOSO
    assert puntuar(uno(objeto="Servicio de comedor escolar en fiestas", cpv="55520000"), config) is None
    assert puntuar(uno(objeto="Paellas para la residencia en fiestas"), config)[0] == ALTO
    assert puntuar(uno(objeto="Suministro de papel", cpv="30190000"), config) is None


def test_zona(config):
    assert en_zona(uno(), ["46"])
    fuera = uno(nuts="ES300", dir3="L01280796", cp="28001", ciudad="Madrid")
    assert not en_zona(fuera, ["46"])
    solo_cp = uno(nuts="", dir3="L01000000", cp="46100")
    assert en_zona(solo_cp, ["46"])
    assert not en_zona(uno(), ["03"])
    alicante = uno(nuts="ES521", dir3="L01030141", cp="03001", ciudad="Alicante")
    castellon = uno(nuts="ES522", dir3="L01120402", cp="12001", ciudad="Castellón")
    murcia = uno(nuts="ES620", dir3="L01300308", cp="30001", ciudad="Murcia")
    toda = uno(nuts="ES52", dir3="A10000000", cp="")
    comunitat = config["provincias"]
    assert comunitat == ["46", "03", "12"]
    assert all(en_zona(c, comunitat) for c in (uno(), alicante, castellon, toda))
    assert not en_zona(murcia, comunitat) and not en_zona(fuera, comunitat)


def test_urls_historico():
    urls = urls_historico("https://x/s/feed.atom", 2, date(2026, 3, 15))
    assert urls == ["https://x/s/feed_2024.zip", "https://x/s/feed_2025.zip",
                    "https://x/s/feed_202601.zip", "https://x/s/feed_202602.zip"]


def test_feed_reciente_sigue_paginas_y_para():
    paginas = {
        "p1": feed(entrada("1", "Paellas fiestas", actualizado="2026-10-31T10:00:00+01:00"),
                   siguiente="p2"),
        "p2": feed(entrada("2", "Paellas Fira", actualizado="2026-10-25T10:00:00+01:00"),
                   entrada("1", "Paellas fiestas (vieja)", actualizado="2026-10-24T10:00:00+01:00"),
                   siguiente="p3"),
        "p3": feed(entrada("3", "Antigua", actualizado="2026-09-01T10:00:00+01:00"), siguiente="p4"),
    }
    pedidas = []

    def obtener(url):
        pedidas.append(url)
        return paginas[url]

    contratos, n = leer_feed_reciente("p1", datetime(2026, 10, 20), "licitaciones", obtener)
    assert pedidas == ["p1", "p2", "p3"]  # p3 ya es anterior: no pide p4
    assert n == 3
    por_id = {c.id[-1]: c for c in contratos}
    assert set(por_id) == {"1", "2"}
    assert por_id["1"].objeto == "Paellas fiestas"  # se queda la versión más reciente


def test_compradores(tmp_path, config):
    datos = feed(
        entrada("1", "Paellas fiestas 2024", estado="RES", adjudicatario="Paellas Pepe SL",
                adjudicado="8000", fecha_adj="2024-07-10"),
        entrada("2", "Paellas fiestas 2025", estado="RES", adjudicatario="Arrocería Sol SL",
                adjudicado="9000", fecha_adj="2025-07-15"),
        entrada("3", "Catering fiestas", estado="RES", organo="Ajuntament d'Alzira", ciudad="Alzira",
                dir3="L01460176", cp="46600", adjudicatario="Paellas Pepe SL", adjudicado="5000",
                fecha_adj="2025-08-20"),
        entrada("4", "Paellas en Madrid", nuts="ES300", dir3="L01280796", cp="28001", ciudad="Madrid"),
        entrada("5", "Limpieza viaria", cpv="90610000"),
    )
    archivo = tmp_path / "contratosMenores_2025.zip"
    with zipfile.ZipFile(archivo, "w") as z:
        z.writestr("a.atom", datos)
    hallazgos = compradores.filtrar(leer_archivo(archivo, "menores"), config)
    assert len(hallazgos) == 3
    municipios = compradores.resumen_municipios(hallazgos, 2)
    sueca, alzira = municipios
    assert sueca["municipio"] == "Sueca" and sueca["prioridad_sugerida"] == 1
    assert sueca["contratos"] == 2 and sueca["enviar_propuesta_en"] == "may"
    assert sueca["dir3_ayto"] == "L01462357"
    assert alzira["municipio"] == "Alzira" and alzira["prioridad_sugerida"] == 2
    competencia = compradores.resumen_competencia(hallazgos)
    assert competencia[0]["empresa"] == "Paellas Pepe SL" and competencia[0]["contratos"] == 2
    res = compradores.guardar(hallazgos, 2, tmp_path / "compradores.xlsx")
    assert res == {"contratos": 3, "municipios": 2, "prioridad_1": 1}


def test_licitaciones_conserva_lo_escrito_a_mano(tmp_path, config):
    ruta = tmp_path / "licitaciones.xlsx"
    contratos = list(leer_atom(feed(
        entrada("1", "Paellas gigantes Fiestas Mayores"),
        entrada("2", "Paellas fiestas (plazo vencido)", plazo="2026-10-01"),
        entrada("3", "Paellas fiestas adjudicada", estado="ADJ"),
        entrada("4", "Servicio de catering"),
    )))
    encontradas = licitaciones.abiertas(contratos, config, AHORA)
    assert {i[-1] for i in encontradas} == {"1", "4"}

    existentes = licitaciones.cargar(ruta)
    nuevas = licitaciones.actualizar(existentes, encontradas, AHORA)
    assert len(nuevas) == 2
    licitaciones.guardar(existentes, AHORA, ruta)

    # El usuario marca una como ME INTERESA y escribe una nota
    from openpyxl import load_workbook
    libro = load_workbook(ruta)
    hoja = libro.active
    cab = [c.value for c in hoja[1]]
    for fila in hoja.iter_rows(min_row=2):
        if str(fila[cab.index("id")].value).endswith("/1"):
            fila[cab.index("seguimiento")].value = "ME INTERESA"
            fila[cab.index("notas")].value = "pedir pliegos"
    libro.save(ruta)

    # Segunda búsqueda: cambia el importe
    contratos = list(leer_atom(feed(entrada("1", "Paellas gigantes Fiestas Mayores", importe="15000",
                                            plazo="2026-11-03"))))
    existentes = licitaciones.cargar(ruta)
    nuevas = licitaciones.actualizar(existentes, licitaciones.abiertas(contratos, config, AHORA), AHORA)
    assert nuevas == []
    licitaciones.guardar(existentes, AHORA, ruta)
    existentes = licitaciones.cargar(ruta)
    fila = next(f for i, f in existentes.items() if i.endswith("/1"))
    assert fila["seguimiento"] == "ME INTERESA" and fila["notas"] == "pedir pliegos"
    assert fila["importe"] == 15000 and fila["dias_restantes"] == 2
    pronto = licitaciones.recordatorios(existentes, AHORA, 3)
    assert len(pronto) == 1


def test_comando_buscar_con_archivo(tmp_path, monkeypatch):
    import vigilante
    monkeypatch.setattr(licitaciones, "RUTA", tmp_path / "licitaciones.xlsx")
    monkeypatch.setattr(vigilante, "RUTA_ESTADO", tmp_path / "estado.json")
    monkeypatch.setattr(vigilante, "avisar", lambda *a: None)
    import fichas
    monkeypatch.setattr(fichas, "CARPETA", tmp_path / "fichas")
    archivo = tmp_path / "licitaciones.atom"
    archivo.write_bytes(feed(entrada("1", "Paellas fiestas", plazo="2099-01-01")))
    assert vigilante.main(["buscar", "--archivo", str(archivo)]) == 0
    assert (tmp_path / "licitaciones.xlsx").exists()
    carpetas = list((tmp_path / "fichas").iterdir())
    assert len(carpetas) == 1 and (carpetas[0] / "00_RESUMEN.txt").exists()
    fila = next(iter(licitaciones.cargar(tmp_path / "licitaciones.xlsx").values()))
    assert fila["valoracion"] > 0 and fila["carpeta"] == str(carpetas[0])


def test_una_fuente_rota_no_para_la_busqueda(tmp_path, monkeypatch):
    """La plataforma devuelve una página HTML en una fuente: se sigue con la otra."""
    import json

    import fichas
    import placsp
    import vigilante
    monkeypatch.setattr(licitaciones, "RUTA", tmp_path / "licitaciones.xlsx")
    monkeypatch.setattr(vigilante, "RUTA_ESTADO", tmp_path / "estado.json")
    monkeypatch.setattr(vigilante, "DATOS", tmp_path)
    monkeypatch.setattr(fichas, "CARPETA", tmp_path / "fichas")
    avisos = []
    monkeypatch.setattr(vigilante, "avisar", lambda t, m: avisos.append(m))

    buena = feed(entrada("1", "Paellas fiestas", plazo="2099-01-01",
                         actualizado=datetime.now().isoformat(timespec="seconds")))
    html = b'<!DOCTYPE html><html><head><meta charset="utf-8"><title>Error</title></head><body><p>x</body></html>'

    class Respuesta:
        def __init__(self, contenido):
            self.content = contenido
            self.status_code = 200
            self.headers = {}

        def raise_for_status(self):
            pass

    def falso_get(self, url, **kw):
        return Respuesta(html if "1044" in url else buena)

    monkeypatch.setattr(placsp.Red, "_get", falso_get)
    monkeypatch.setattr(placsp.Red, "descargar_documento", lambda self, url, c, n: (_ for _ in ()).throw(OSError()))
    assert vigilante.main(["buscar"]) == 0
    assert len(licitaciones.cargar(tmp_path / "licitaciones.xlsx")) == 1
    estado = json.loads((tmp_path / "estado.json").read_text())["ultima_busqueda"]
    assert "licitaciones" in estado and "agregadas" not in estado
    assert any("agregadas" in a for a in avisos)
